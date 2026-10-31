import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from pathlib import Path

repo_root = Path(r"c:\Users\root\Desktop\_Intelligence\nullspread-research-academy")
docx_path = repo_root / "knowledge_base" / "domains" / "computer_science" / "referat_binary_numeral_systems.docx"
brain_dir = Path(r"C:\Users\root\.gemini\antigravity\brain\2af54222-aa64-4815-b67b-8de6182563f5")

doc = docx.Document()

# 1. Page Setup according to Odessa Polytechnic regulations
section = doc.sections[0]
section.page_width = Inches(210 / 25.4)   # A4 width: 210mm
section.page_height = Inches(297 / 25.4)  # A4 height: 297mm
section.top_margin = Inches(20 / 25.4)     # 20mm
section.bottom_margin = Inches(20 / 25.4)  # 20mm
section.left_margin = Inches(25 / 25.4)    # 25mm (standard binding)
section.right_margin = Inches(20 / 25.4)   # 20mm (balanced)

# Enable different first page for title page
section.different_first_page_header_footer = True

# Add page number field to regular header (top-right)
header = section.header
header_p = header.paragraphs[0]
header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header_run = header_p.add_run()
header_run.font.name = "Times New Roman"
header_run.font.size = Pt(12)

# Insert Word PAGE field into header
fldSimple = OxmlElement('w:fldSimple')
fldSimple.set(qn('w:instr'), 'PAGE')
header_p._p.append(fldSimple)

# Set base styles
normal_style = doc.styles['Normal']
normal_style.font.name = 'Times New Roman'
normal_style.font.size = Pt(14)
normal_style.font.color.rgb = RGBColor(0, 0, 0)
normal_style.paragraph_format.line_spacing = 1.5
normal_style.paragraph_format.space_after = Pt(0)
normal_style.paragraph_format.space_before = Pt(0)

def add_title_page():
    # University header
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_uni.paragraph_format.line_spacing = 1.15
    p_uni.paragraph_format.space_after = Pt(4)
    r1 = p_uni.add_run("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ\n")
    r1.bold = True
    r1.font.size = Pt(12.5)
    r2 = p_uni.add_run("НАЦІОНАЛЬНИЙ УНІВЕРСИТЕТ «ОДЕСЬКА ПОЛІТЕХНІКА»\n")
    r2.bold = True
    r2.font.size = Pt(13)
    r3 = p_uni.add_run("Інститут комп'ютерних систем\nКафедра комп'ютерних інтелектуальних систем та мереж")
    r3.font.size = Pt(12)
    
    # Space before middle
    p_space1 = doc.add_paragraph()
    p_space1.paragraph_format.space_before = Pt(72)
    
    # Doc type & topic
    p_mid = doc.add_paragraph()
    p_mid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mid.paragraph_format.line_spacing = 1.3
    
    r_type = p_mid.add_run("РЕФЕРАТ\n")
    r_type.bold = True
    r_type.font.size = Pt(16)
    
    r_disc = p_mid.add_run("з дисципліни: «Комп'ютерна схемотехніка та архітектура комп'ютерів»\n\n")
    r_disc.font.size = Pt(13.5)
    
    r_top_lbl = p_mid.add_run("на тему:\n")
    r_top_lbl.font.size = Pt(13)
    
    r_topic = p_mid.add_run("Системи числення, двійкова арифметика та принципи машинного представлення даних в обчислювальних системах\n")
    r_topic.bold = True
    r_topic.font.size = Pt(14)
    
    # Space before student block
    p_space2 = doc.add_paragraph()
    p_space2.paragraph_format.space_before = Pt(50)
    
    # Student / Teacher metadata block (aligned right, left-aligned text)
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.left_indent = Inches(3.2)
    p_meta.paragraph_format.line_spacing = 1.25
    p_meta.paragraph_format.space_after = Pt(0)
    
    r_st = p_meta.add_run("Виконав:\nстудент 1 курсу, групи АМ-251\nденної форми навчання\nРудік Л. О.\n\n")
    r_st.font.size = Pt(13)
    
    r_tch = p_meta.add_run("Перевірив:\nдоцент кафедри КІСМ, к.т.н.\nКузнєцов М. О.\n")
    r_tch.font.size = Pt(13)
    
    # Space before footer
    p_space3 = doc.add_paragraph()
    p_space3.paragraph_format.space_before = Pt(80)
    
    # Footer City & Year
    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot = p_foot.add_run("Одеса – 2026")
    r_foot.bold = True
    r_foot.font.size = Pt(13)
    
    # Page break to TOC
    doc.add_page_break()

def add_toc():
    # TOC Header centered
    p_toc_head = doc.add_paragraph()
    p_toc_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_toc_head.paragraph_format.space_after = Pt(18)
    r = p_toc_head.add_run("ЗМІСТ")
    r.bold = True
    r.font.size = Pt(14)
    
    toc_items = [
        ("ВСТУП", "3", True, 0),
        ("РОЗДІЛ 1. ЗАГАЛЬНІ ВІДОМОСТІ ПРО СИСТЕМИ ЧИСЛЕННЯ", "4", True, 0),
        ("1.1. Поняття системи числення. Непозиційні та позиційні системи", "4", False, 0.3),
        ("1.2. Базові характеристики позиційних систем числення", "4", False, 0.3),
        ("РОЗДІЛ 2. ОСНОВНІ СИСТЕМИ ЧИСЛЕННЯ В ОБЧИСЛЮВАЛЬНІЙ ТЕХНІЦІ", "6", True, 0),
        ("2.1. Двійкова система числення як основа цифрової електроніки", "6", False, 0.3),
        ("2.2. Вісімкова та шістнадцяткова системи числення", "7", False, 0.3),
        ("РОЗДІЛ 3. АЛГОРИТМИ ПЕРЕВЕДЕННЯ ЧИСЕЛ ТА ДВІЙКОВА АРИФМЕТИКА", "8", True, 0),
        ("3.1. Переведення цілих та дробових чисел між системами числення", "8", False, 0.3),
        ("3.2. Швидкий перехід між двійковою, вісімковою та шістнадцятковою системами", "9", False, 0.3),
        ("3.3. Двійкова арифметика та машинне кодування чисел зі знаком", "10", False, 0.3),
        ("ВИСНОВКИ", "12", True, 0),
        ("СПИСОК ВИКОРИСТАНИХ ДЖЕРЕЛ", "13", True, 0),
    ]
    
    # Printable width = 210mm - 25mm - 20mm = 165mm = 6.496 inches
    tab_stop_pos = Inches(6.45)
    
    for title, page_str, is_bold, indent_in in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(4)
        if indent_in > 0:
            p.paragraph_format.left_indent = Inches(indent_in)
            
        # Add right-aligned tab with dot leader
        p.paragraph_format.tab_stops.add_tab_stop(tab_stop_pos, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        
        run_title = p.add_run(title)
        run_title.bold = is_bold
        run_title.font.size = Pt(13 if not is_bold else 13.5)
        
        run_tab = p.add_run(f"\t{page_str}")
        run_tab.bold = is_bold
        run_tab.font.size = Pt(13 if not is_bold else 13.5)
        
    doc.add_page_break()

def add_major_heading(title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(14)

def add_sub_heading(title):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(1.25 / 2.54)
    p.paragraph_format.left_indent = Inches(0)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(14)

def add_p(text, bold_prefix="", indent=True, keep_with_next=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = keep_with_next
    if indent:
        p.paragraph_format.first_line_indent = Inches(1.25 / 2.54)
        p.paragraph_format.left_indent = Inches(0)
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.bold = True
    r = p.add_run(text)
    return p

def add_list_item(num_str, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(1.75 / 2.54)
    p.paragraph_format.first_line_indent = Inches(-0.5 / 2.54)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.35
    r_num = p.add_run(num_str + "\t")
    r_num.bold = True
    p.add_run(text)

def add_formula(runs):
    # runs is a list of tuples: (text, is_bold, is_sub, is_sup)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    for text, bold, sub, sup in runs:
        r = p.add_run(text)
        r.bold = bold
        r.font.subscript = sub
        r.font.superscript = sup
        r.font.size = Pt(14)
    return p

# 2. Build Document Body
add_title_page()
add_toc()

# ВСТУП
add_major_heading("ВСТУП")
add_p(
    "Сучасна цифрова техніка та мікропроцесорні пристрої працюють з електричними сигналами дискретної природи. Всі дані, що зберігаються в пам'яті ЕОМ та обробляються центральним процесором (числа, тексти, команди, графічні зображення), на фізичному рівні подаються у вигляді двійкових кодів. Вибір способів кодування та систем числення визначається законами напівпровідникової схемотехніки та дискретної математики.",
    bold_prefix="Актуальність теми. "
)
add_p(
    "Використання десяткової системи в електронних схемах призвело б до надмірного ускладнення апаратури: для надійного розпізнавання десяти різних рівнів напруги потрібні дуже точні й дорогі компоненти, чутливі до перешкод. Натомість двійкова система спирається лише на два дискретні стани сигналу (низький та високий рівні напруги), що забезпечує максимальну завадостійкість і простоту побудови логічних елементів. Для скороченого запису двійкових кодів фахівці використовують вісімкову та шістнадцяткову системи числення."
)
add_p(
    "вивчити теоретичні засади систем числення, з'ясувати причини застосування двійкової, вісімкової та шістнадцяткової систем в обчислювальній техніці, опанувати правила переведення чисел і виконання двійкових арифметичних операцій, а також розглянути машинне представлення цілих чисел зі знаком у прямому, оберненому та додатковому кодах.",
    bold_prefix="Мета реферату — "
)
add_p("Основні завдання:", indent=True, keep_with_next=True)
add_list_item("1.", "Розглянути відмінності між непозиційними та позиційними системами числення.")
add_list_item("2.", "Дослідити розгорнуту степеневу форму запису чисел.")
add_list_item("3.", "З'ясувати фізичні та логічні переваги двійкової системи в схемотехніці ЕОМ.")
add_list_item("4.", "Ознайомитися з призначенням вісімкової та шістнадцяткової систем.")
add_list_item("5.", "Розглянути практичні алгоритми переведення цілих і дробових чисел.")
add_list_item("6.", "Опанувати двійкове додавання, множення та принцип заміни віднімання додаванням у додатковому коді.")

# РОЗДІЛ 1
doc.add_page_break()
add_major_heading("РОЗДІЛ 1. ЗАГАЛЬНІ ВІДОМОСТІ ПРО СИСТЕМИ ЧИСЛЕННЯ")
add_sub_heading("1.1. Поняття системи числення. Непозиційні та позиційні системи")
add_p("Системою числення називають сукупність правил і символів, за допомогою яких записуються, читаються та обробляються числа [2]. Символи, що використовуються для запису чисел, називають цифрами, а повний набір цих цифр утворює алфавіт системи числення.")
add_p("Усі відомі системи числення прийнято поділяти на два основні класи: непозиційні та позиційні.")
add_p("В непозиційних системах числення кількісне значення цифри є постійним і не залежить від того, яку позицію вона посідає в числі. Історично першою була одинична система (палички, зарубки на дереві чи кістках). Найвідомішим прикладом розвиненої непозиційної системи є римська нумерація, в якій використовуються літери латинського алфавіту: I = 1, V = 5, X = 10, L = 50, C = 100, D = 500, M = 1000. Значення числа визначається сумуванням або відніманням значень цифр (наприклад, VI = 5 + 1 = 6, а IV = 5 - 1 = 4).")
add_p("Головними недоліками непозиційних систем є:", indent=True, keep_with_next=True)
add_list_item("1.", "Відсутність цифри нуль для позначення порожнього розряду.")
add_list_item("2.", "Неможливість зручного запису дробових або великих чисел.")
add_list_item("3.", "Надзвичайна складність виконання арифметичних дій (множення або ділення чисел у римському записі є громіздким процесом).")
add_p("У позиційних системах числення значення цифри визначається її позицією (розрядом) у записі числа. Позиційний принцип був винайдений у стародавньому Вавилоні та Індії, після чого поширився світом завдяки працям арабських математиків. Прикладом є звична для нас десяткова система. Наприклад, у числі 555 крайня права цифра означає п'ять простих одиниць, середня — п'ять десятків, а ліва — п'ять сотень.")

add_sub_heading("1.2. Базові характеристики позиційних систем числення")
add_p("Будь-яка позиційна система однозначно описується кількома основними параметрами:", indent=True, keep_with_next=True)
add_list_item("1.", "Основа системи числення (b) — ціле додатне число (b ≥ 2), яке дорівнює кількості символів в алфавіті цієї системи. Основа показує, у скільки разів одиниця наступного (старшого) розряду більша за одиницю поточного.")
add_list_item("2.", "Алфавіт системи числення — упорядкований набір цифр. Якщо основа не перевищує десяти (b ≤ 10), застосовують арабські цифри від 0 до b - 1. Якщо основа більша за десять, алфавіт доповнюють літерами латинського алфавіту (наприклад, A = 10, B = 11 тощо).")
add_list_item("3.", "Вага розряду — степінь основи b^k, де k — номер позиції цифри відносно коми (ціла частина нумерується від нуля праворуч наліво, дробова — від мінус одиниці ліворуч направо).")

add_p("У довільній системі числення з основою b будь-яке дійсне число N можна записати у вигляді суми степенів основи [1]:")
add_formula([
    ("N = a", False, False, False),
    ("n-1", False, True, False),
    ("b", False, False, False),
    ("n-1", False, False, True),
    (" + a", False, False, False),
    ("n-2", False, True, False),
    ("b", False, False, False),
    ("n-2", False, False, True),
    (" + ... + a", False, False, False),
    ("0", False, True, False),
    ("b", False, False, False),
    ("0", False, False, True),
    (" + a", False, False, False),
    ("-1", False, True, False),
    ("b", False, False, False),
    ("-1", False, False, True),
    (" + ... + a", False, False, False),
    ("-m", False, True, False),
    ("b", False, False, False),
    ("-m", False, False, True),
])
add_p("де a_k — відповідні цифри з алфавіту системи числення; n — кількість цілих розрядів; m — кількість дробових розрядів.")
add_p("Наприклад, розгорнутий запис десяткового числа 348,25 має вигляд:")
add_formula([
    ("348,25", False, False, False),
    ("10", False, True, False),
    (" = 3·10", False, False, False),
    ("2", False, False, True),
    (" + 4·10", False, False, False),
    ("1", False, False, True),
    (" + 8·10", False, False, False),
    ("0", False, False, True),
    (" + 2·10", False, False, False),
    ("-1", False, False, True),
    (" + 5·10", False, False, False),
    ("-2", False, False, True),
])
add_formula([
    ("= 300 + 40 + 8 + 0,2 + 0,05 = 348,25", False, False, False)
])
add_p("Ця розгорнута формула лежить в основі універсального правила переведення чисел з будь-якої системи числення в десяткову.")

# РОЗДІЛ 2
doc.add_page_break()
add_major_heading("РОЗДІЛ 2. ОСНОВНІ СИСТЕМИ ЧИСЛЕННЯ В ОБЧИСЛЮВАЛЬНІЙ ТЕХНІЦІ")
add_sub_heading("2.1. Двійкова система числення як основа цифрової електроніки")
add_p("Двійкова система числення має основу b = 2 та містить лише дві цифри: 0 і 1. Один двійковий розряд прийнято називати бітом.")
add_p("Вибір двійкової системи для побудови цифрової апаратури зумовлений такими важливими факторами [4]:", indent=True, keep_with_next=True)
add_list_item("1.", "Простота та надійність електронних компонентів. Електронні транзистори найзручніше використовувати як двопозиційні ключі: закритий стан транзистора (струм не проходить, напруга низька) відповідає логічному нулю, а відкритий стан (струм проходить, напруга висока) — логічній одиниці. Завдяки наявності чіткої різниці між рівнями напруги схеми мають високу завадостійкість.")
add_list_item("2.", "Простий зв'язок з математичною логікою. Двійкові цифри 0 та 1 точно відповідають логічним значенням «хибність» (False) та «істина» (True). Усі операції над ними реалізуються елементарними логічними вентилями (І, АБО, НЕ), що дозволяє створювати складні арифметичні вузли процесора (суматори, компаратори).")
add_list_item("3.", "Максимальна простота правил арифметики. Таблиці додавання та множення для одного двійкового розряду є елементарними.")

# Add Word Tables for Truth Tables
p_tbl1_lbl = doc.add_paragraph()
p_tbl1_lbl.paragraph_format.left_indent = Inches(1.25 / 2.54)
p_tbl1_lbl.paragraph_format.space_before = Pt(8)
p_tbl1_lbl.paragraph_format.space_after = Pt(2)
p_tbl1_lbl.paragraph_format.keep_with_next = True
r = p_tbl1_lbl.add_run("Таблиця 2.1. Двійкове додавання")
r.italic = True
r.font.size = Pt(12)

tbl1 = doc.add_table(rows=3, cols=3)
tbl1.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl1_data = [
    ["+", "0", "1"],
    ["0", "0", "1"],
    ["1", "1", "10₂"]
]
for r_idx, row in enumerate(tbl1.rows):
    for c_idx, cell in enumerate(row.cells):
        cell.text = tbl1_data[r_idx][c_idx]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].runs[0].font.name = "Times New Roman"
        cell.paragraphs[0].runs[0].font.size = Pt(12)
        if r_idx == 0 or c_idx == 0:
            cell.paragraphs[0].runs[0].bold = True

p_tbl2_lbl = doc.add_paragraph()
p_tbl2_lbl.paragraph_format.left_indent = Inches(1.25 / 2.54)
p_tbl2_lbl.paragraph_format.space_before = Pt(8)
p_tbl2_lbl.paragraph_format.space_after = Pt(2)
p_tbl2_lbl.paragraph_format.keep_with_next = True
r2 = p_tbl2_lbl.add_run("Таблиця 2.2. Двійкове множення")
r2.italic = True
r2.font.size = Pt(12)

tbl2 = doc.add_table(rows=3, cols=3)
tbl2.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl2_data = [
    ["×", "0", "1"],
    ["0", "0", "0"],
    ["1", "0", "1"]
]
for r_idx, row in enumerate(tbl2.rows):
    for c_idx, cell in enumerate(row.cells):
        cell.text = tbl2_data[r_idx][c_idx]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].runs[0].font.name = "Times New Roman"
        cell.paragraphs[0].runs[0].font.size = Pt(12)
        if r_idx == 0 or c_idx == 0:
            cell.paragraphs[0].runs[0].bold = True

add_p("Слід звернути увагу на зв'язок двійкової арифметики з логічними функціями: результат додавання двох бітів (молодший розряд суми) формується операцією «Виключне АБО» (XOR), а біт переносу до старшого розряду — операцією кон'юнкції «І» (AND). Двійкове ж множення однорозрядних чисел повністю збігається з логічною операцією «І».")

add_sub_heading("2.2. Вісімкова та шістнадцяткова системи числення")
p_hex = doc.add_paragraph()
p_hex.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_hex.paragraph_format.left_indent = Inches(1.25 / 2.54)
p_hex.add_run("Для людини двійкові записи є незручними, оскільки навіть невелике число або адреса пам'яті вимагає довгої послідовності нулів та одиниць. Наприклад, число 48879")
r_sub1 = p_hex.add_run("10")
r_sub1.font.subscript = True
p_hex.add_run(" у двійковій системі записується як 1011111011101111")
r_sub2 = p_hex.add_run("2")
r_sub2.font.subscript = True
p_hex.add_run(". Такий запис важко сприймати та легко припуститися помилки під час запису.")

add_p("Для компактного представлення двійкової інформації використовують системи числення з основами, які є степенями двійки: вісімкову (8 = 2^3) та шістнадцяткову (16 = 2^4) [6].")
add_p("В алфавіті шістнадцяткової системи використовуються цифри від 0 до 9 та перші шість літер латинського алфавіту:")
add_formula([
    ("A = 10,   B = 11,   C = 12,   D = 13,   E = 14,   F = 15", False, False, False)
])
add_p("Головна перевага шістнадцяткової системи полягає в тому, що кожна шістнадцяткова цифра замінює рівно 4 двійкових розряди (тетраду). Відповідно, один машинний байт (8 бітів) компактно записується рівно двома шістнадцятковими знаками — від 00 до FF.")

p_beef = doc.add_paragraph()
p_beef.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_beef.paragraph_format.left_indent = Inches(1.25 / 2.54)
p_beef.add_run("Наведене вище 16-розрядне число 1011111011101111")
r = p_beef.add_run("2")
r.font.subscript = True
p_beef.add_run(" у шістнадцятковій формі записується коротко: BEEF")
r = p_beef.add_run("16")
r.font.subscript = True
p_beef.add_run(".")

add_p("Основні галузі практичного застосування шістнадцяткової системи:", indent=True, keep_with_next=True)
add_list_item("1.", "Запис адрес оперативної пам'яті та регістрів процесора.")
add_list_item("2.", "Подання кольорів у комп'ютерній графіці та веб-дизайні (формат RGB, де кожна складова записується двома шістнадцятковими цифрами, наприклад #FF8000).")
add_list_item("3.", "Запис фізичних адрес мережевих адаптерів (MAC-адреси).")
add_list_item("4.", "Перегляд вмісту двійкових файлів і дампів пам'яті.")
add_p("Вісімкова система замінює трійки бітів (тріади) однією вісімковою цифрою (від 0 до 7). Вона застосовувалася в ранніх обчислювальних машинах, а сьогодні зустрічається, наприклад, при налаштуванні прав доступу до файлів в операційній системі Linux.")

# РОЗДІЛ 3
doc.add_page_break()
add_major_heading("РОЗДІЛ 3. АЛГОРИТМИ ПЕРЕВЕДЕННЯ ЧИСЕЛ ТА ДВІЙКОВА АРИФМЕТИКА")
add_sub_heading("3.1. Переведення цілих та дробових чисел між системами числення")
add_p("Для переведення чисел з довільної системи в десяткову використовують розгорнуту степеневу форму запису.")

p_ex1 = doc.add_paragraph()
p_ex1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_ex1.paragraph_format.left_indent = Inches(1.25 / 2.54)
p_ex1.paragraph_format.keep_with_next = True
r_bld = p_ex1.add_run("Приклад 1. ")
r_bld.bold = True
p_ex1.add_run("Перевести двійкове число 110101")
r = p_ex1.add_run("2")
r.font.subscript = True
p_ex1.add_run(" у десяткову систему:")

add_formula([
    ("110101", False, False, False),
    ("2", False, True, False),
    (" = 1·2", False, False, False),
    ("5", False, False, True),
    (" + 1·2", False, False, False),
    ("4", False, False, True),
    (" + 0·2", False, False, False),
    ("3", False, False, True),
    (" + 1·2", False, False, False),
    ("2", False, False, True),
    (" + 0·2", False, False, False),
    ("1", False, False, True),
    (" + 1·2", False, False, False),
    ("0", False, False, True),
])
add_formula([
    ("= 32 + 16 + 0 + 4 + 0 + 1 = 53", False, False, False),
    ("10", False, True, False),
])

p_ex2 = doc.add_paragraph()
p_ex2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_ex2.paragraph_format.left_indent = Inches(1.25 / 2.54)
p_ex2.paragraph_format.keep_with_next = True
r_bld = p_ex2.add_run("Приклад 2. ")
r_bld.bold = True
p_ex2.add_run("Перевести шістнадцяткове число 2A3")
r = p_ex2.add_run("16")
r.font.subscript = True
p_ex2.add_run(" у десяткову систему:")

add_formula([
    ("2A3", False, False, False),
    ("16", False, True, False),
    (" = 2·16", False, False, False),
    ("2", False, False, True),
    (" + 10·16", False, False, False),
    ("1", False, False, True),
    (" + 3·16", False, False, False),
    ("0", False, False, True),
    (" = 2·256 + 160 + 3 = 675", False, False, False),
    ("10", False, True, False),
])

add_p("Для переведення цілого десяткового числа в двійкову чи іншу систему застосовують алгоритм послідовного ділення на нову основу [3]:", indent=True, keep_with_next=True)
add_list_item("1.", "Число ділять націло на основу нової системи (наприклад, на 2).")
add_list_item("2.", "Залишок від ділення записують як чергову молодшу цифру.")
add_list_item("3.", "Отриману частку знову ділять на основу.")
add_list_item("4.", "Процес продовжують доти, доки частка не стане рівною нулю.")
add_list_item("5.", "Результат записують із залишків у зворотному порядку (від останнього до першого).")

p_ex3 = doc.add_paragraph()
p_ex3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_ex3.paragraph_format.left_indent = Inches(1.25 / 2.54)
p_ex3.paragraph_format.keep_with_next = True
r_bld = p_ex3.add_run("Приклад 3. ")
r_bld.bold = True
p_ex3.add_run("Перевести число 29")
r = p_ex3.add_run("10")
r.font.subscript = True
p_ex3.add_run(" у двійкову систему:")

for step_txt in [
    "29 / 2 = 14 (залишок 1)",
    "14 / 2 = 7 (залишок 0)",
    "7 / 2 = 3 (залишок 1)",
    "3 / 2 = 1 (залишок 1)",
    "1 / 2 = 0 (залишок 1)"
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(1.6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    p.add_run("• " + step_txt)

p_res3 = doc.add_paragraph()
p_res3.paragraph_format.left_indent = Inches(1.25 / 2.54)
p_res3.add_run("Записуючи залишки знизу вгору, отримуємо: 29")
r = p_res3.add_run("10")
r.font.subscript = True
p_res3.add_run(" = 11101")
r = p_res3.add_run("2")
r.font.subscript = True
p_res3.add_run(".")

add_p("Для переведення правильних дробів застосовують алгоритм послідовного множення на нову основу:", indent=True, keep_with_next=True)
add_list_item("1.", "Дробову частину множать на основу нової системи.")
add_list_item("2.", "Ціла частина отриманого добутку стає черговою цифрою дробу.")
add_list_item("3.", "Дробову частину знову множать на основу.")
add_list_item("4.", "Процес зупиняють при отриманні нульової дробової частини або після досягнення потрібної кількості знаків.")
add_list_item("5.", "Отримані цілі частини виписують у прямому порядку (згори вниз).")

p_ex4 = doc.add_paragraph()
p_ex4.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_ex4.paragraph_format.left_indent = Inches(1.25 / 2.54)
p_ex4.paragraph_format.keep_with_next = True
r_bld = p_ex4.add_run("Приклад 4. ")
r_bld.bold = True
p_ex4.add_run("Перевести десятковий дріб 0,375")
r = p_ex4.add_run("10")
r.font.subscript = True
p_ex4.add_run(" у двійковий вигляд:")

for step_txt in [
    "0,375 × 2 = 0,75 (ціла частина 0)",
    "0,75 × 2 = 1,5 (ціла частина 1, залишається дріб 0,5)",
    "0,5 × 2 = 1,0 (ціла частина 1, залишок 0)"
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(1.6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    p.add_run("• " + step_txt)

p_res4 = doc.add_paragraph()
p_res4.paragraph_format.left_indent = Inches(1.25 / 2.54)
p_res4.add_run("Отримуємо результат: 0,375")
r = p_res4.add_run("10")
r.font.subscript = True
p_res4.add_run(" = 0,011")
r = p_res4.add_run("2")
r.font.subscript = True
p_res4.add_run(".")

add_p("Слід підкреслити важливу особливість: звичайний десятковий дріб, наприклад 0,1 або 0,2, у двійковій системі перетворюється на нескінченний періодичний дріб:")
add_formula([
    ("0,1", False, False, False),
    ("10", False, True, False),
    (" ≈ 0,0001100110011...2 = 0,0(0011)", False, False, False),
    ("2", False, True, False)
])
add_p("Оскільки в пам'яті комп'ютера розрядна сітка завжди обмежена фіксованою кількістю бітів, нескінченний дріб округлюється. Саме через це при роботі з дійсними числами в комп'ютерних програмах виникають невеликі похибки обчислень (наприклад, результат 0,1 + 0,2 дає приблизно 0,30000000000000004).")

add_sub_heading("3.2. Швидкий перехід між двійковою, вісімковою та шістнадцятковою системами")
add_p("Оскільки основи 8 = 2^3 та 16 = 2^4 є цілими степенями двійки, перехід між ними виконується без проміжних обчислень через десяткову систему — за допомогою таблиць відповідності груп бітів (тріад і тетрад).")
add_p("Для переведення з двійкової системи у шістнадцяткову:", indent=True, keep_with_next=True)
add_list_item("1.", "Двійкове число розбивають на групи по 4 цифри (тетради) в обидва боки від коми (цілу частину — праворуч наліво, дробову — ліворуч направо).")
add_list_item("2.", "Якщо в крайніх групах не вистачає розрядів, їх доповнюють нулями.")
add_list_item("3.", "Кожну тетраду замінюють однією відповідною шістнадцятковою цифрою.")

p_ex5 = doc.add_paragraph()
p_ex5.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_ex5.paragraph_format.left_indent = Inches(1.25 / 2.54)
p_ex5.paragraph_format.keep_with_next = True
r_bld = p_ex5.add_run("Приклад 5. ")
r_bld.bold = True
p_ex5.add_run("Перевести число 110101101")
r = p_ex5.add_run("2")
r.font.subscript = True
p_ex5.add_run(" у шістнадцяткову систему:")

for t_item in [
    "1101 → D",
    "1010 → A",
    "1 → доповнюємо нулями зліва 0001 → 1"
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(1.6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    p.add_run("• " + t_item)

p_res5 = doc.add_paragraph()
p_res5.paragraph_format.left_indent = Inches(1.25 / 2.54)
p_res5.add_run("Отримуємо: 110101101")
r = p_res5.add_run("2")
r.font.subscript = True
p_res5.add_run(" = 1AD")
r = p_res5.add_run("16")
r.font.subscript = True
p_res5.add_run(".")
add_p("Зворотний перехід здійснюється ще простіше: кожна шістнадцяткова цифра безпосередньо розгортається у чотири біти.")

add_sub_heading("3.3. Двійкова арифметика та машинне кодування чисел зі знаком")
add_p("Додавання багаторозрядних двійкових чисел виконується у стовпчик за тими самими правилами, що й у десятковій системі, але з основою 2:")

p_add_col = doc.add_paragraph()
p_add_col.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_add_col.paragraph_format.line_spacing = 1.2
p_add_col.paragraph_format.space_before = Pt(6)
p_add_col.paragraph_format.space_after = Pt(6)
p_add_col.paragraph_format.keep_with_next = True
p_add_col.add_run("   101101₂   (45₁₀)\n+ 011011₂   (27₁₀)\n─────────────\n   1001000₂   (72₁₀)")

add_p("Перевірка показує: 45 + 27 = 72, що підтверджує правильність розрахунку.")
add_p("Множення у двійковій системі виконується зсувом та додаванням. Множення на одиницю дає саме множене, а множення на нуль дає нулі, що значно спрощує схемну реалізацію множника.")

add_p("Кодування від'ємних чисел у пам'яті ЕОМ.", indent=True, keep_with_next=True)
add_p("У комп'ютерній техніці для представлення від'ємних цілих чисел використовують старший розряд розрядної сітки: 0 означає додатне число, а 1 — від'ємне. Відомо три основні способи кодування [3, 5]:", indent=True, keep_with_next=True)

add_list_item("1.", "Прямий код. У старшому розряді зберігається знак, а в решті — двійковий модуль числа. Наприклад, для 8-бітної сітки: +5₁₀ = 00000101₂, -5₁₀ = 10000101₂. Недоліками прямого коду є наявність двох нулів (+0 та -0), а також необхідність створення окремого складного блоку віднімання.")
add_list_item("2.", "Обернений код. Для від'ємних чисел усі біти модуля інвертуються (замінюються на протилежні): +5₁₀ = 00000101₂, -5₁₀ = 11111010₂. Тут також залишається проблема двох нулів і виникає потреба в циклічному переносі одиниці.")
add_list_item("3.", "Додатковий код. Це сучасний загальноприйнятий стандарт в мікропроцесорній техніці. Для отримання додаткового коду від'ємного числа потрібно:")

for sub_step in [
    "записати число в прямому коді;",
    "інвертувати всі його інформаційні розряди (отримати обернений код);",
    "додати двійкову одиницю до молодшого розряду."
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(1.8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    p.add_run("• " + sub_step)

p_ex6 = doc.add_paragraph()
p_ex6.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_ex6.paragraph_format.left_indent = Inches(1.25 / 2.54)
p_ex6.paragraph_format.keep_with_next = True
r_bld = p_ex6.add_run("Приклад 6. ")
r_bld.bold = True
p_ex6.add_run("Записати число -5")
r = p_ex6.add_run("10")
r.font.subscript = True
p_ex6.add_run(" у 8-бітному додатковому коді:")

for s in [
    "1. Число +5₁₀ у 8 бітах: 00000101",
    "2. Інверсія бітів: 11111010",
    "3. Додавання одиниці: 11111010₂ + 1 = 11111011₂"
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(1.6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.keep_with_next = True
    p.add_run(s)

p_res6 = doc.add_paragraph()
p_res6.paragraph_format.left_indent = Inches(1.25 / 2.54)
p_res6.paragraph_format.keep_with_next = True
p_res6.add_run("Отримуємо: [-5]")
r = p_res6.add_run("дод")
r.font.subscript = True
p_res6.add_run(" = 11111011₂.")

add_p("Головні переваги додаткового коду:", indent=True, keep_with_next=True)
add_list_item("1.", "Єдиний нуль. Число нуль завжди має однозначний вигляд 00000000. При спробі закодувати «мінус нуль» одиниця переносу виходить за межі 8 розрядів і скидається процесором.")

# The subtraction example must NEVER be split across pages
p_sub_hdr = doc.add_paragraph()
p_sub_hdr.paragraph_format.left_indent = Inches(1.25 / 2.54)
p_sub_hdr.paragraph_format.keep_with_next = True
r_num2 = p_sub_hdr.add_run("2. Заміна віднімання додаванням. ")
r_num2.bold = True
p_sub_hdr.add_run("Операція A - B обчислюється як A + [-B]")
r = p_sub_hdr.add_run("дод")
r.font.subscript = True
p_sub_hdr.add_run(". Завдяки цьому процесору не потрібен окремий апаратний блок віднімання — всі операції виконує один швидкий суматор.")

p_sub_ex = doc.add_paragraph()
p_sub_ex.paragraph_format.left_indent = Inches(1.25 / 2.54)
p_sub_ex.paragraph_format.keep_with_next = True
p_sub_ex.add_run("Наприклад, віднімання 12 - 5:")

for sub_line in [
    "+12₁₀ = 00001100₂",
    "-5₁₀ = 11111011₂ (у додатковому коді)",
    "Сума: 00001100₂ + 11111011₂ = 1 00000111₂"
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(1.6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.keep_with_next = True
    p.add_run("• " + sub_line)

p_sub_conc = doc.add_paragraph()
p_sub_conc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_sub_conc.paragraph_format.left_indent = Inches(1.25 / 2.54)
p_sub_conc.add_run("Старша одиниця виходить за межі 8-бітної сітки й відкидається. Залишається результат: 00000111₂ = +7₁₀, що є абсолютно правильним.")

p_over = doc.add_paragraph()
p_over.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_over.paragraph_format.left_indent = Inches(1.25 / 2.54)
r_over_b = p_over.add_run("Поняття переповнення розрядної сітки. ")
r_over_b.bold = True
p_over.add_run("Якщо результат арифметичної операції перевищує допустимий діапазон для поточної кількості бітів (для 8 бітів зі знаком це значення від -128 до +127), біти результату спотворюють знаковий розряд. У таких випадках процесор встановлює спеціальний прапорець переповнення (Overflow Flag), що сигналізує про помилку в розрахунках.")

# ВИСНОВКИ
doc.add_page_break()
add_major_heading("ВИСНОВКИ")
add_p("У ході виконання реферату було досліджено теоретичні основи та практичні засади функціонування систем числення в обчислювальних машинах:", indent=True, keep_with_next=True)
add_list_item("1.", "Доведено, що позиційні системи числення мають абсолютні переваги над непозиційними завдяки наявності нуля, компактності запису та простоті виконання арифметичних операцій за допомогою степеневого розкладу.")
add_list_item("2.", "З'ясовано, що домінування двійкової системи в схемотехніці зумовлене надійністю дворівневих напівпровідникових елементів, стійкістю до перешкод та простотою реалізації логічних операцій.")
add_list_item("3.", "Встановлено, що вісімкова та шістнадцяткова системи відіграють роль зручного проміжного подання машинних двійкових кодів для людини, забезпечуючи швидкий потетрадний перехід без виконання складних ділень.")
add_list_item("4.", "Розглянуто алгоритми переводу цілих і дробових чисел та з'ясовано причину виникнення похибок при збереженні десяткових дробів у скінченній розрядній сітці.")
add_list_item("5.", "Проаналізовано принципи машинного кодування чисел зі знаком та показано, чому додатковий код став основним стандартом: він вирішує проблему однозначності нуля та дозволяє апаратно реалізувати операцію віднімання за допомогою звичайного двійкового суматора.")

# СПИСОК ВИКОРИСТАНИХ ДЖЕРЕЛ
doc.add_page_break()
add_major_heading("СПИСОК ВИКОРИСТАНИХ ДЖЕРЕЛ")
sources = [
    "Бабич М. П., Жуков І. А. Комп'ютерна схемотехніка : навч. посіб. Київ : КНТ, 2008. 488 с.",
    "ДСТУ ISO/IEC 2382:2015. Інформаційні технології. Словник термінів. Київ : ДП «УкрНДНЦ», 2016. 344 с.",
    "Матвієнко М. П. Архітектура комп’ютера : підручник. 2-ге вид., переробл. і доповн. Київ : Ліра-К, 2021. 544 с.",
    "Мельник А. О. Архітектура комп'ютерів : підручник. Львів : Видавництво Національного університету «Львівська політехніка», 2008. 470 с.",
    "Таненбаум Е., Остін Т. Архітектура комп'ютера. 6-те вид. Бостон : Pearson, 2013. 800 с.",
    "Харіс Д. М., Харіс С. Л. Цифрова схемотехніка та архітектура комп'ютера. 2-ге вид. Волтгем : Morgan Kaufmann, 2012. 712 с."
]
for idx, s in enumerate(sources, 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(1.25 / 2.54)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.35
    r_n = p.add_run(f"{idx}. ")
    r_n.bold = True
    p.add_run(s)

doc.save(str(docx_path))
import shutil
shutil.copy2(str(docx_path), str(brain_dir / "referat_binary_numeral_systems.docx"))
print("Successfully generated perfect DOCX document!")
